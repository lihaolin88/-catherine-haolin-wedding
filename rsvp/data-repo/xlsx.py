"""Turn RSVP responses (a list of dicts) into a formatted Excel sheet."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# (key in the submitted JSON, column header in the sheet, column width)
COLUMNS = [
    ("submittedAt", "Submitted (UTC)", 20),
    ("name", "Name", 18),
    ("arrival", "Arrival in Hawaii", 16),
    ("firstCourse", "First Course", 30),
    ("mainCourse", "Main Course", 30),
    ("allergies", "Allergies", 24),
    ("drinks", "Alcohol / Wine", 24),
    ("notes", "Anything We Should Know", 36),
    ("wishes", "Blessings", 48),
    ("lang", "Lang", 6),
]


def write_xlsx(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "RSVP"
    ws.append([header for _, header, _ in COLUMNS])
    for row in rows:
        ws.append([row.get(key, "") for key, _, _ in COLUMNS])

    header_fill = PatternFill("solid", fgColor="A3823A")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    # Dinner head-count so the caterer numbers are one glance away.
    summary = wb.create_sheet("Dinner Count")
    summary.append(["Course", "Choice", "Count"])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for key, label in (("firstCourse", "First Course"), ("mainCourse", "Main Course")):
        counts = {}
        for row in rows:
            choice = row.get(key) or "(none)"
            counts[choice] = counts.get(choice, 0) + 1
        for choice, n in sorted(counts.items()):
            summary.append([label, choice, n])
    summary.append([])
    summary.append(["Total responses", "", len(rows)])
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 36
    summary.column_dimensions["C"].width = 8

    wb.save(path)
